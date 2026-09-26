import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path

import openpyxl

from app.core.models import Check, Project
from app.core.normalize import excel_date, norm_key
from app.core.overlap import center, distance_mi

PROJECT_HEADERS = ["project_id", "utility", "state", "project_name", "name_a", "lat_a", "lon_a", "name_b", "lat_b",
                   "lon_b", "lat_center", "lon_center", "in_service_date", "overlap_count", "overlap_1", "overlap_2",
                   "overlap_3"]
OVERLAP_HEADERS = ["overlap_id", "distance_mi", "time_gap (day)", "utility_a", "project_id_a", "project_name_a",
                   "utility_b", "project_id_b", "project_name_b"]


@dataclass
class SamplePoint:
    name: str
    lat: float | None
    lon: float | None


@dataclass
class SampleProject:
    ref_id: str  # 'DESC_3'
    utility: str
    name: str
    a: SamplePoint
    b: SamplePoint
    in_service: date
    in_service_is_text: bool

    @property
    def center(self) -> tuple[float, float] | None:
        return center([(self.a.lat, self.a.lon), (self.b.lat, self.b.lon)])


@dataclass
class SampleOverlap:
    overlap_id: str
    a: str
    b: str
    distance_mi: float
    time_gap_days: int


@dataclass
class Sample:
    projects: dict[str, SampleProject]
    overlaps: list[SampleOverlap]


def _num(v: object) -> float | None:
    return float(v) if isinstance(v, (int, float)) else None


def read_sample(path: Path) -> Sample:
    wb = openpyxl.load_workbook(path, data_only=False)
    rows = list(wb["projects"].iter_rows(values_only=True))
    if list(rows[0][: len(PROJECT_HEADERS)]) != PROJECT_HEADERS:
        raise ValueError("Projects_Overlaps.xlsx 'projects' headers changed")
    projects: dict[str, SampleProject] = {}
    for r in rows[1:]:
        if not r[0]:
            continue
        rec = dict(zip(PROJECT_HEADERS, r))
        d = excel_date(rec["in_service_date"])
        if d is None:
            raise ValueError(f"{rec['project_id']}: unreadable in_service_date {rec['in_service_date']!r}")
        projects[rec["project_id"]] = SampleProject(
            ref_id=rec["project_id"], utility=rec["utility"], name=rec["project_name"],
            a=SamplePoint(rec["name_a"] or "", _num(rec["lat_a"]), _num(rec["lon_a"])),
            b=SamplePoint(rec["name_b"] or "", _num(rec["lat_b"]), _num(rec["lon_b"])),
            in_service=d, in_service_is_text=isinstance(rec["in_service_date"], str),
        )
    orows = list(wb["overlaps"].iter_rows(values_only=True))
    if list(orows[0][: len(OVERLAP_HEADERS)]) != OVERLAP_HEADERS:
        raise ValueError("Projects_Overlaps.xlsx 'overlaps' headers changed")
    overlaps = [SampleOverlap(overlap_id=o[0], a=o[4], b=o[7], distance_mi=float(o[1]), time_gap_days=int(o[2]))
                for o in orows[1:] if o[0]]
    return Sample(projects=projects, overlaps=overlaps)


def _key(name: str) -> str:
    return re.sub(r"[^a-z0-9]", "", name.lower())


def match_projects(sample: Sample, projects: list[Project]) -> dict[str, str]:
    # Match by name; break ties with the in-service date.
    out: dict[str, str] = {}
    for ref_id, sp in sample.projects.items():
        utility = "DESC" if ref_id.startswith("DESC") else "GA"
        cands = [p for p in projects if p.utility == utility and _key(p.name) == _key(sp.name)]
        if len(cands) > 1:
            cands = [p for p in cands if p.in_service_date == sp.in_service.isoformat()] or cands[:1]
        if cands:
            out[ref_id] = cands[0].id
    return out


def sample_checks(sample: Sample) -> list[Check]:
    # Problems inside the sponsor's own file.
    checks: list[Check] = []
    src = "Projects_Overlaps.xlsx"
    by_name: dict[str, set[tuple[float, float, str]]] = {}
    for sp in sample.projects.values():
        for pt in (sp.a, sp.b):
            if pt.name and pt.lat is not None and pt.lon is not None:
                by_name.setdefault(norm_key(pt.name), set()).add((round(pt.lat, 6), round(pt.lon, 6), sp.ref_id))
    for key, pts in by_name.items():
        coords = sorted({(la, lo) for la, lo, _ in pts})
        if len(coords) > 1:
            gap = distance_mi(coords[0], coords[1])
            refs = ", ".join(sorted({r for _, _, r in pts}))
            checks.append(Check(id=f"dupcoord:{key}", level="warn", rule="same_name_two_coords",
                                title="Same substation, two locations in the benchmark file",
                                detail=f"{key.title()} appears with {len(coords)} different coordinates ({refs}), about "
                                       f"{gap:.2f} mi apart. Each project keeps its own row's coordinates, so the "
                                       "reference distances still match.", source=src))
    text = [sp.ref_id for sp in sample.projects.values() if sp.in_service_is_text]
    real = [sp.ref_id for sp in sample.projects.values() if not sp.in_service_is_text]
    if text and real:
        checks.append(Check(id="sample:mixed_dates", level="warn", rule="mixed_date_types",
                            title="Mixed date formats in the benchmark file",
                            detail=f"Most in-service dates are text like '12/31/2024', but {', '.join(real)} "
                                   "are stored as real Excel dates. Both were normalized before comparing.", source=src))
    missing = [f"{pt.name} ({sp.ref_id})" for sp in sample.projects.values() for pt in (sp.a, sp.b)
               if pt.name and pt.lat is None]
    if missing:
        checks.append(Check(id="sample:missing_coords", level="info", rule="sample_missing_coords",
                            title="Benchmark endpoints without coordinates",
                            detail=f"{', '.join(missing)} have no coordinates, so those projects use their one "
                                   "located endpoint as the center.", source=src))
    return checks
