# Sperry's columns first, in their order. Our extra columns go after.

import io
from datetime import date

import openpyxl
from openpyxl.styles import Font

from app.core.models import Overlap, Project
from app.core.owners import Book, book
from app.core.sample import OVERLAP_HEADERS, PROJECT_HEADERS
from app.store import dataset

# Owner names and codes come from the sources table. sponsor is the source's code (DESC, GPC, ...);
# filing_sponsor keeps the owner the filing itself names for the row (GPC, SAV, GTC, MEAG, DU).
PROJECT_EXTRA = ["sponsor", "location_confidence", "build_start", "estimated_cost", "status", "source", "project_type",
                 "date_precision", "filing_sponsor"]
OVERLAP_EXTRA = ["location_confidence", "windows_overlap", "in_sponsor_sample", "center_distance_mi", "tier",
                 "other_utilities_nearby"]
CATEGORY_NAME = {"electric": "electric", "gas": "gas", "roads_water": "roads and water"}


def safe_cell(v: object) -> object:
    # Text that starts like a formula is stored as text, so a submitted name can't run in Excel or Sheets.
    return f"'{v}" if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r") else v


def _append(ws, row: list) -> None:
    ws.append([safe_cell(v) for v in row])


def owner_label(p: Project, owners: Book | None = None) -> str:
    return (owners or book()).display_name(p)


def _project_row(p: Project, overlaps_of: dict[str, list[str]], owners: Book) -> list:
    title = [e for e in p.endpoints if e.role == "endpoint"]  # description places aren't endpoints
    a = title[0] if title else None
    b = title[1] if len(title) > 1 else None
    ids = overlaps_of.get(p.id, [])
    return [p.id, owners.display_name(p), owners.state_of(p), p.name,
            a.name if a else None, a.lat if a else None, a.lon if a else None,
            b.name if b else None, b.lat if b else None, b.lon if b else None,
            p.lat, p.lon, date.fromisoformat(p.in_service_date), len(ids),
            ids[0] if ids else None, ids[1] if len(ids) > 1 else None, ids[2] if len(ids) > 2 else None,
            owners.code(p), p.location_confidence, date.fromisoformat(p.build_start) if p.build_start else None,
            p.cost_total, p.status,
            f"{p.source_file} p.{p.source_page} ({p.source_ref})" if owners.builtin(p) else f"{p.source_file}, {p.source_ref}",
            p.project_type, p.date_precision or "day", p.sponsor]


def _header(ws, headers: list[str]) -> None:
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)


def build_xlsx(overlaps: list[Overlap]) -> bytes:
    ids = {o.id: f"OVL_{i}" for i, o in enumerate(overlaps, start=1)}
    overlaps_of: dict[str, list[str]] = {}
    for o in overlaps:
        overlaps_of.setdefault(o.project_a, []).append(ids[o.id])
        overlaps_of.setdefault(o.project_b, []).append(ids[o.id])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "projects"
    owners = book()
    _header(ws, PROJECT_HEADERS + PROJECT_EXTRA)
    for p in sorted(dataset.CURRENT.projects.values(), key=lambda p: (p.utility, p.id)):
        _append(ws, _project_row(p, overlaps_of, owners))
    for row in ws.iter_rows(min_row=2):
        for idx in (12, 19):  # in_service_date, build_start
            if row[idx].value:
                row[idx].number_format = "mm/dd/yyyy"
    links = dataset.others(overlaps)
    research = {r.id: r for r in dataset.CURRENT.research}
    near: dict[str, list[str]] = {}
    for t in links:
        r = research[t.research_id]
        near.setdefault(t.overlap_id, []).append(f"{r.utility}: {r.name} ({t.dist_a_mi} / {t.dist_b_mi} mi)")
    ws2 = wb.create_sheet("overlaps")
    _header(ws2, OVERLAP_HEADERS + OVERLAP_EXTRA)
    for o in overlaps:
        a, b = dataset.CURRENT.projects[o.project_a], dataset.CURRENT.projects[o.project_b]
        _append(ws2, [ids[o.id], o.distance_mi, o.time_gap_days, owners.display_name(a), a.id, a.name,
                    owners.display_name(b), b.id, b.name, o.pair_confidence, o.windows_overlap, o.in_sponsor_sample, o.center_mi, o.tier,
                    "; ".join(near.get(o.id, []))])
    ws3 = wb.create_sheet("data_checks")
    _header(ws3, ["level", "rule", "title", "detail", "source", "project_id", "decided_by"])
    for c in dataset.CURRENT.checks:
        _append(ws3, [c.level, c.rule, c.title, c.detail, c.source, c.project_id, c.actor])
    ws4 = wb.create_sheet("reference_test")
    _header(ws4, ["overlap_id", "sponsor_a", "sponsor_b", "our_a", "our_b", "expected_mi", "got_mi",
                  "expected_days", "got_days", "pass"])
    for r in dataset.CURRENT.reference:
        _append(ws4, [r.overlap_id, r.a, r.b, r.a_project, r.b_project, r.expected_mi, r.got_mi, r.expected_days,
                    r.got_days, r.passed])
    ws6 = wb.create_sheet("other_utilities")
    _header(ws6, ["research_id", "category", "owner", "project", "status", "start_as_stated", "in_service_as_stated",
                  "in_service_date_used", "location_confidence", "lat", "lon", "places", "sources", "found_by",
                  "fact_checkers_confirmed"])
    for r in dataset.CURRENT.research:
        v = r.verification or {}
        _append(ws6, [r.id, CATEGORY_NAME[r.category], r.utility, r.name, r.status, r.start, r.in_service, r.in_service_date,
                    r.location_confidence, r.lat, r.lon, "; ".join(f"{p.name} ({p.kind}, {p.state})" for p in r.places),
                    " ".join(s.url for s in r.sources), r.found_by,
                    f"{v.get('confirmed')} of {v.get('verifiers')}" if v.get("verifiers") else None])
    ws7 = wb.create_sheet("other_utility_links")
    _header(ws7, ["overlap_id", "research_id", "owner", "project", "miles_to_project_a", "miles_to_project_b",
                  "days_from_project_a_in_service", "days_from_project_b_in_service", "date_is_year_or_month_only"])
    for t in links:
        r = research[t.research_id]
        _append(ws7, [ids.get(t.overlap_id, t.overlap_id), r.id, r.utility, r.name, t.dist_a_mi, t.dist_b_mi,
                    t.gap_a_days, t.gap_b_days, t.approx_date])
    ws5 = wb.create_sheet("notes")
    for line in [
        "distance_mi: miles between the closest points of the two projects (the challenge's rule). A line is the "
        "straight segment between its two located ends; anything else is a point.",
        "center_distance_mi: center to center, as in Sperry's sample (center = midpoint of located endpoints, or the "
        "single located one; haversine miles, R = 3958.8). The reference_test sheet uses this one.",
        "Overlap: closest points under 25 miles apart. tier: touching (under 0.1 mi), row (under 1 mi), site (under "
        "5 mi), crew (under 25 mi). time_gap (day) = |in-service date A - in-service date B|.",
        "Dates: DESC 'Planned In-Service Date' (final phase if phased); Georgia 'Need Date' from each project page.",
        "Georgia sponsor scope in this export follows the filters used when exporting.",
        "location_confidence: verified (sponsor file / override), confirmed_osm (OpenStreetMap + judge), "
        "partial (mixed), town (GeoNames town, approximate), unlocated.",
        "Georgia costs are REDACTED in the filing and exported as empty, never zero.",
        "Other utilities: projects found by the research team, each with cited sources. A project is listed on an "
        "overlap when its center is under 25 miles from BOTH project centers. Year-only dates use Dec 31.",
        f"Dataset run: {dataset.CURRENT.run_id}",
    ]:
        ws5.append([line])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
