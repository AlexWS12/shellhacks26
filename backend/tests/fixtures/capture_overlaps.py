# Records the overlap output of the offline pipeline, so tests/test_snapshot.py can check a refactor changed nothing.
# Run from backend/:  uv run python tests/fixtures/capture_overlaps.py   (only when a change to the output is intended)

import asyncio
import datetime
import io
import json
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

import openpyxl  # noqa: E402

from app import config  # noqa: E402

OUT = Path(__file__).with_name("overlaps_snapshot.json")
VIEWS = {  # the filter combinations the UI and exports use
    "default": {}, "all_sponsors": {"all_sponsors": True}, "verified_only": {"min_confidence": "confirmed_osm"},
    "hide_finished": {"hide_finished": True}, "by_gap": {"sort": "gap"},
}


def cell(v):
    return v.isoformat() if isinstance(v, (datetime.date, datetime.datetime)) else v


def capture() -> dict:
    from app.core.overlap import Filters
    from app.core.sample import OVERLAP_HEADERS, PROJECT_HEADERS
    from app.export import build_xlsx
    from app.pipeline import build_pipeline
    from app.runtime.executor import execute
    from app.runtime.run import new_run
    from app.store import dataset

    config.GEMINI_API_KEY, config.JEV_PROVIDER, config.OSM_LIVE = "", "", False
    config.SUBMISSIONS_DIR = Path(tempfile.mkdtemp())
    config.SOURCES_DB = Path(tempfile.mkdtemp()) / "sources.db"  # a freshly seeded table: DESC and GPC only
    run = new_run("live")
    run.pace = 0
    asyncio.run(execute(run, *build_pipeline()))
    board = run.board
    dataset.load_snapshot(board.to_snapshot())
    out = {"pipeline": [o.model_dump() for o in board.overlaps],
           "projects": {p.id: [p.lat, p.lon, p.location_confidence] for p in board.projects.values()},
           "views": {k: [o.model_dump() for o in dataset.overlaps(Filters(today=config.TODAY, **v))] for k, v in VIEWS.items()}}
    wb = openpyxl.load_workbook(io.BytesIO(build_xlsx(dataset.overlaps(Filters(today=config.TODAY)))))
    out["export_projects"] = [[cell(c.value) for c in row[: len(PROJECT_HEADERS)]] for row in wb["projects"].iter_rows()]
    out["export_overlaps"] = [[cell(c.value) for c in row[: len(OVERLAP_HEADERS)]] for row in wb["overlaps"].iter_rows()]
    return out


if __name__ == "__main__":
    OUT.write_text(json.dumps(capture(), indent=0, default=str), encoding="utf-8")
    print(f"wrote {OUT}")
