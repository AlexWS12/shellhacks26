# The research team: loading, dates, placement, the category choice, the three-way check, live-search citations.
import asyncio
import io
import json

import openpyxl
import pytest

from app import config
from app.agents.research import live_records
from app.clients.gemini import cite
from app.core.models import Overlap, Project, ResearchProject
from app.core.research import load_file, nearby, normalize_date, prepare
from app.runtime.executor import execute
from app.runtime.run import new_run

SRC = [{"url": "https://example.org/plan", "title": "Plan", "publisher": "Test", "quote": "q"}]
RECORDS = [
    {"category": "electric", "utility": "Test Co-op", "name": "Jasper County Test Substation", "status": "planned",
     "start": "2026", "in_service": "2027-06", "places": [{"name": "Jasper County", "kind": "county", "state": "SC", "role": "site"}],
     "sources": SRC},
    {"category": "gas", "utility": "Test Pipeline", "name": "Chatham Test Lateral", "in_service": "2028",
     "places": [{"name": "Chatham County", "kind": "county", "state": "GA", "role": "site"}], "sources": SRC},
    {"category": "roads_water", "utility": "Test DOT", "name": "Road With No Town", "in_service": None,
     "places": [{"name": "I-95", "kind": "road", "state": "SC", "role": "along"}], "sources": SRC},
    {"category": "electric", "utility": "No Source Co", "name": "Dropped Without Sources", "places": [], "sources": []},
]


@pytest.fixture(scope="module")
def research_file(tmp_path_factory):
    p = tmp_path_factory.mktemp("research") / "other_utilities.json"
    p.write_text(json.dumps({"generated": "2026-09-26", "records": RECORDS}), encoding="utf-8")
    return p


def run_with(research_file, categories):
    from app.pipeline import build_pipeline

    old = config.RESEARCH_FILE
    config.GEMINI_API_KEY, config.JEV_PROVIDER, config.OSM_LIVE, config.RESEARCH_LIVE = "", "", False, False
    config.RESEARCH_FILE = research_file
    try:
        run = new_run("live")
        run.pace = 0
        run.research = categories
        asyncio.run(execute(run, *build_pipeline()))
        return run
    finally:
        config.RESEARCH_FILE = old


@pytest.fixture(scope="module")
def all_run(research_file):
    return run_with(research_file, ["electric", "gas", "roads_water"])


def test_dates_keep_their_precision():
    assert normalize_date("2028", end=True) == ("2028-12-31", "year")
    assert normalize_date("2028", end=False) == ("2028-01-01", "year")
    assert normalize_date("2027-02", end=True) == ("2027-02-28", "month")
    assert normalize_date("2026-10-01", end=True) == ("2026-10-01", "day")
    assert normalize_date("late 2027", end=True) == (None, None)
    assert normalize_date(None, end=True) == (None, None)


def test_records_without_sources_are_dropped(research_file):
    records, meta = load_file(research_file)
    assert meta["generated"] == "2026-09-26"
    assert [r.name for r in records] == ["Jasper County Test Substation", "Chatham Test Lateral", "Road With No Town"]
    assert records[0].in_service_date == "2027-06-30" and records[0].date_precision == "month"


def test_run_places_records_and_explains_the_rest(all_run):
    last = all_run.events[-1]
    assert last["type"] == "run.done" and last["ok"], last
    r = {x.name: x for x in all_run.board.research.values()}
    jasper = r["Jasper County Test Substation"]
    assert jasper.lat is not None and jasper.endpoints[0].method == "county_centroid" and jasper.location_confidence == "town"
    road = r["Road With No Town"]
    assert road.lat is None and "roads are not placed" in road.endpoints[0].evidence["reason"]
    types = [e["type"] for e in all_run.events]
    assert types.count("research.found") == 3 and "research.placed" in types and "research.unlocated" in types
    started = next(e for e in all_run.events if e["type"] == "run.started")
    assert started["research"] == ["electric", "gas", "roads_water"]


def test_three_way_links_are_code_and_under_25_miles_from_both(all_run):
    b = all_run.board
    assert b.third_party, "expected the Jasper County record near a Savannah-area opportunity"
    for t in b.third_party:
        o = next(o for o in b.overlaps if o.id == t.overlap_id)
        assert t.dist_a_mi < 25 and t.dist_b_mi < 25
        r = b.research[t.research_id]
        if r.in_service_date:
            from datetime import date
            assert t.gap_a_days == abs((date.fromisoformat(r.in_service_date) - date.fromisoformat(b.projects[o.project_a].in_service_date)).days)
            assert t.approx_date == (r.date_precision in ("year", "month"))
    assert len([e for e in all_run.events if e["type"] == "third_party.found"]) == len(b.third_party)


def test_only_chosen_categories_run(research_file):
    run = run_with(research_file, ["electric"])
    assert {r.category for r in run.board.research.values()} == {"electric"}
    gas = next(e for e in run.events if e["type"] == "agent.done" and e["agent_id"] == "research_gas")
    assert gas["summary"] == "not selected for this run"
    assert all(t.category == "electric" for t in run.board.third_party)
    assert run.board.research_selected == ["electric"]


def test_benchmark_is_unchanged_by_research(all_run):
    assert len(all_run.board.reference) == 6 and all(r.passed for r in all_run.board.reference)


def test_nearby_needs_both_sides_under_25_miles():
    def proj(pid, lat, lon, util):
        return Project(id=pid, utility=util, sponsor=util if util == "DESC" else "GPC", name=pid, in_service_date="2027-12-31",
                       source_file="x", source_page=1, source_ref=pid, lat=lat, lon=lon, location_confidence="verified")
    a, b = proj("A", 32.30, -81.00, "DESC"), proj("B", 32.30, -81.30, "GA")
    o = Overlap(id="A|B", project_a="A", project_b="B", distance_mi=17.6, time_gap_days=0, pair_confidence="verified")
    mk = lambda rid, lat, lon: ResearchProject(id=rid, category="gas", utility="U", name=rid, lat=lat, lon=lon,  # noqa: E731
                                               location_confidence="town", in_service_date="2028-12-31", date_precision="year")
    between, far_side = mk("between", 32.30, -81.15), mk("far", 32.30, -80.75)  # far: ~15 mi from A, ~32 from B
    links = nearby([o], {"A": a, "B": b}, [between, far_side])
    assert [t.research_id for t in links] == ["between"]
    assert links[0].gap_a_days == 366 and links[0].approx_date and links[0].confidence == "town"
    assert nearby([o], {"A": a, "B": b}, [between], categories=["electric"]) == []


def test_live_search_keeps_only_cited_records():
    sources = [{"url": "https://a.example/1", "title": "a.example"}, {"url": "https://b.example/2", "title": "b.example"}]
    raw = [{"utility": "X", "name": "Cited Project", "description": "", "status": "planned", "start": None, "in_service": "2029",
            "date_quote": None, "places": [], "source_numbers": [2]},
           {"utility": "Y", "name": "Uncited Project", "description": "", "status": "planned", "start": None,
            "in_service": None, "date_quote": None, "places": [], "source_numbers": [7]}]
    known = [prepare({"category": "gas", "utility": "Z", "name": "Already Known", "sources": SRC})]
    raw.append({**raw[0], "name": "Already Known"})
    out = live_records(raw, sources, "gas", known)
    assert [r.name for r in out] == ["Cited Project"]
    assert out[0].sources[0].url == "https://b.example/2" and out[0].found_by == "gemini_search" and out[0].id.startswith("LIVE-")
    assert cite("Alpha. Beta.", [{"end": 6, "chunks": [1]}, {"end": 12, "chunks": [0, 1]}]) == "Alpha.[2] Beta.[1][2]"


def test_export_and_api_include_other_utilities(all_run):
    from fastapi.testclient import TestClient

    from app.export import build_xlsx
    from app.main import app
    from app.store import dataset

    dataset.load_snapshot(all_run.board.to_snapshot())
    wb = openpyxl.load_workbook(io.BytesIO(build_xlsx(all_run.board.overlaps)))
    assert wb["overlaps"][1][-1].value == "other_utilities_nearby"
    assert wb["other_utilities"].max_row == 1 + len(all_run.board.research)
    assert wb["other_utility_links"].max_row == 1 + len(all_run.board.third_party)
    body = TestClient(app).get("/api/research").json()
    assert body["selected"] == ["electric", "gas", "roads_water"]
    assert len(body["records"]) == 3 and len(body["links"]) == len(all_run.board.third_party)


def test_malformed_records_are_skipped_not_fatal(tmp_path):
    p = tmp_path / "bad.json"
    p.write_text(json.dumps({"records": [RECORDS[0], {**RECORDS[1], "miles": "about ten"}, {"category": "gas"}]}), encoding="utf-8")
    records, meta = load_file(p)
    assert [r.name for r in records] == ["Jasper County Test Substation"] and meta["skipped"] == 2


def test_county_center_only_when_nothing_finer_is_known():
    from app.core.models import Endpoint, ResearchPlace
    from app.core.research import place_center

    town = Endpoint(name="Hardeeville", lat=32.28, lon=-81.08, method="geonames_town", confidence="town")
    county = Endpoint(name="Jasper County", lat=32.44, lon=-81.03, method="county_centroid", confidence="town")
    site = lambda n: ResearchPlace(name=n, kind="town", state="SC", role="site")  # noqa: E731
    r = ResearchProject(id="x", category="electric", utility="U", name="x", places=[site("Hardeeville"), site("Jasper County")],
                        endpoints=[town, county])
    assert place_center(r) == (32.28, -81.08)
    along = [ResearchPlace(name=n, kind="county", state="GA", role="along") for n in ("A County", "B County")]
    r2 = ResearchProject(id="y", category="gas", utility="U", name="y", places=along, endpoints=[
        Endpoint(name="A County", lat=33.0, lon=-82.0, method="county_centroid", confidence="town"),
        Endpoint(name="B County", lat=33.2, lon=-82.4, method="county_centroid", confidence="town")])
    assert place_center(r2) == (33.1, -82.2)  # a pipeline known only by the counties it crosses


def test_approx_filter_also_drops_approximate_links(all_run):
    from app.store import dataset

    dataset.load_snapshot(all_run.board.to_snapshot())
    everything = dataset.others(all_run.board.overlaps, "town")
    precise = dataset.others(all_run.board.overlaps, "confirmed_osm")
    assert everything and all(t.confidence == "town" for t in everything) and precise == []
