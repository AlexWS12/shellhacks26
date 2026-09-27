# Utilities as data: the sources table, what reads it (Reader dispatch, OSM search area and operator names, the
# overlap engine, exports, /api/sources) and a third source pairing with both built-ins.

import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import config
from app.core import overlap
from app.core.models import Endpoint, Project
from app.core.overlap import Filters, find_overlaps
from app.core.owners import book
from app.core.palette import contrast, delta_e, readable
from app.store import sources, submissions

CSV = (b"Project Name,Project ID,In-Service Date,Latitude,Longitude,State\n"
       b"Test Tie A,T-1,12/31/2028,33.30,-81.60,SC\nTest Tie B,T-2,06/30/2029,32.20,-81.20,GA\n")
MAPPING = {"name": "Project Name", "id": "Project ID", "in_service": "In-Service Date", "lat": "Latitude",
           "lon": "Longitude", "state": "State"}


@pytest.fixture
def db(tmp_path, monkeypatch):
    # A freshly seeded table and no saved plans.
    monkeypatch.setattr(config, "SOURCES_DB", tmp_path / "sources.db")
    monkeypatch.setattr(config, "SUBMISSIONS_DIR", tmp_path / "subs")
    monkeypatch.setattr(config, "OSM_BBOX", "")
    return tmp_path


def test_dominion_and_georgia_are_seeded_as_built_in_active_sources(db):
    desc, gpc = sources.all_sources()
    assert (desc.id, desc.code, desc.reader, desc.status, desc.utility_key, desc.states) == (
        "desc", "DESC", "builtin:desc", "active", "DESC", ["SC"])
    assert (gpc.id, gpc.code, gpc.reader, gpc.status, gpc.utility_key, gpc.states) == (
        "gpc", "GPC", "builtin:gpc", "active", "GA", ["GA"])
    assert "sce&g" in desc.osm_operator_patterns and "south carolina electric" in desc.osm_operator_patterns
    assert gpc.default_sponsors == {"GPC", "SAV"}
    assert desc.file_sha256 and len(desc.file_sha256) == 64  # the PDF the parser reads
    assert (desc.color, gpc.color) == ("#2dd4bf", "#6ea8fe")


def test_the_seed_runs_once_and_edits_stick(db):
    sources.set_status("gpc", "failed")
    sources._ready.clear()  # as after a restart
    assert sources.get("gpc").status == "failed"


def test_search_area_is_the_union_of_active_states_with_an_override(db, monkeypatch):
    assert book().bbox() == (30.3, -85.7, 35.3, -78.4)  # SC + GA, the box the geocoder always used
    sources.add("Duke Energy Progress", "Duke Energy Progress", ["NC"], "ai", status="active")
    s, w, n, e = book().bbox()
    assert n > 36.5 and e > -75.5 and (s, w) == (30.3, -85.7)
    sources.set_status(sources.all_sources()[-1].id, "draft")
    assert book().bbox() == (30.3, -85.7, 35.3, -78.4)  # only active sources count
    monkeypatch.setattr(config, "OSM_BBOX", "31,-84,34,-80")
    assert book().bbox() == (31.0, -84.0, 34.0, -80.0)


def test_osm_operator_names_come_from_each_projects_source(db):
    b = book()
    desc = Project(id="DESC-1", utility="DESC", sponsor="DESC", name="x", in_service_date="2027-01-01",
                   source_file="f", source_page=1, source_ref="r")
    ga = desc.model_copy(update={"id": "GA-1", "utility": "GA", "sponsor": "SAV"})
    assert "sce&g" in b.operators(desc, "SC") and "georgia power" in b.operators(ga, "GA")
    # a source with no names of its own gets those of the active sources in its state
    s = sources.add("Test Co", "Test Co", ["SC"], "ai", status="active")
    mine = desc.model_copy(update={"id": "T-1", "utility": s.utility_key, "source_id": s.id, "sponsor": "Test Co"})
    assert book().operators(mine, "SC") == tuple(sources.get("desc").osm_operator_patterns)


def test_new_sources_get_readable_colors_distinct_from_every_other(db):
    added = [sources.add(f"Utility {i}", f"Utility {i}", ["SC"], "ai") for i in range(4)]
    colors = [s.color for s in sources.all_sources()]
    assert len(set(colors)) == len(colors)
    for s in added:
        assert readable(s.color) and contrast(s.color, "#0c0f15") >= 4.5
        assert min(delta_e(s.color, c) for c in colors if c != s.color) > 15
    same_owner = sources.add("Utility 0", "Utility 0 (second plan)", ["SC"], "ai")
    assert same_owner.color == added[0].color  # one owner, one color


def _project(pid: str, source, lat: float, lon: float, date: str = "2028-06-30") -> Project:
    return Project(id=pid, utility=source.utility_key, source_id=source.id, sponsor=source.display_name, name=pid,
                   in_service_date=date, source_file="hand-written", source_page=1, source_ref=pid, lat=lat, lon=lon,
                   location_confidence="verified", endpoints=[Endpoint(name=pid, lat=lat, lon=lon, method="submitted",
                                                                       confidence="verified")])


def test_a_third_source_pairs_with_both_built_ins_and_the_dominion_georgia_pairs_stay(db, finished_run):
    projects = [p.model_copy(deep=True) for p in finished_run.board.projects.values()]
    before = {(o.id, o.distance_mi, o.time_gap_days) for o in find_overlaps(projects, Filters(today=config.TODAY))}
    third = sources.add("Test Transmission", "Test Transmission", ["SC", "GA"], "ai", status="active")
    near_desc = next(p for p in projects if p.utility == "DESC" and p.lat is not None)
    near_ga = next(p for p in projects if p.utility == "GA" and p.sponsor == "GPC" and p.lat is not None)
    mine = [_project("TT-1", third, near_desc.lat + 0.05, near_desc.lon), _project("TT-2", third, near_ga.lat, near_ga.lon + 0.05)]
    got = find_overlaps(projects + mine, Filters(today=config.TODAY))
    pairs = {(book().group(a), book().group(b)) for o in got
             for a, b in [(next(p for p in projects + mine if p.id == o.project_a), next(p for p in projects + mine if p.id == o.project_b))]}
    assert {("DESC", "TEST-TRANSMISSION"), ("GPC", "TEST-TRANSMISSION")} <= pairs
    assert {(o.id, o.distance_mi, o.time_gap_days) for o in got if "TT-" not in o.id} == before  # same math, same pairs
    assert all(not o.project_a.startswith("TT-") for o in got)  # built-ins stay side A
    sources.set_status(third.id, "draft")
    assert {(o.id, o.distance_mi, o.time_gap_days) for o in find_overlaps(projects + mine, Filters(today=config.TODAY))} == before


def test_two_filings_of_one_owner_are_not_paired_with_each_other(db):
    a = sources.add("Same Co", "Same Co", ["SC"], "ai", status="active")
    b = sources.add("Same Co", "Same Co, second plan", ["SC"], "ai", status="active")
    assert a.code == b.code
    got = find_overlaps([_project("S-1", a, 33.0, -81.0), _project("S-2", b, 33.01, -81.0)], Filters(today=config.TODAY))
    assert got == []


def test_sponsors_hidden_by_default_come_from_the_table(db):
    ga = Project(id="GA-9", utility="GA", sponsor="MEAG", name="x", in_service_date="2027-01-01", source_file="f",
                 source_page=1, source_ref="r", lat=33.0, lon=-83.0, location_confidence="verified")
    assert not overlap.visible(ga, Filters()) and overlap.visible(ga, Filters(all_sponsors=True))
    assert overlap.visible(ga.model_copy(update={"sponsor": "SAV"}), Filters())


def test_readers_are_dispatched_from_the_table(db):
    from app.pipeline import build_pipeline

    agents, cards = build_pipeline()
    ids = [a.spec.id for a in agents]
    assert ids[:3] == ["sample", "extract_desc", "extract_ga"] and [c["id"] for c in cards] == ["desc", "gpc", "sample"]
    geo = next(a for a in agents if a.spec.id == "geocoder")
    assert geo.spec.depends_on == ["sample", "extract_desc", "extract_ga"]
    sources.set_status("desc", "draft")
    agents, cards = build_pipeline()
    assert "extract_desc" not in [a.spec.id for a in agents] and [c["id"] for c in cards] == ["gpc", "sample"]
    assert next(a for a in agents if a.spec.id == "geocoder").spec.depends_on == ["sample", "extract_ga"]


def test_saved_plans_are_sources_too(db):
    s = submissions.create("Santee Cooper", None, "SC", "spreadsheet", "plan.csv", CSV, ".csv",
                           columns=list(MAPPING.values()))
    row = sources.get(s.id)
    assert (row.code, row.reader, row.status, row.utility_key, row.states) == ("SANTEE-COOPER", "sheet", "draft",
                                                                             s.owner_key, ["SC"])
    submissions.set_mapping(s.id, MAPPING)
    assert sources.get(s.id).status == "active"
    submissions.delete(s.id)
    assert sources.get(s.id) is None


def test_a_saved_plan_runs_through_the_pipeline_and_pairs_with_both(db):
    from tests.test_submissions import _run

    s = submissions.create("Test Transmission", None, "SC", "spreadsheet", "plan.csv", CSV, ".csv",
                           columns=list(MAPPING.values()))
    submissions.set_mapping(s.id, MAPPING)
    run = _run()
    assert run.events[-1]["ok"], run.events[-1]
    b = run.board
    mine = {p.id for p in b.projects.values() if p.source_id == s.id}
    assert mine and all(b.projects[i].utility == s.owner_key for i in mine)
    sides = {b.projects[o.project_a].utility for o in b.overlaps if o.project_b in mine}
    assert sides == {"DESC", "GA"}


def test_exports_keep_sperry_columns_and_use_the_source_code(db, finished_run):
    from app.core.sample import PROJECT_HEADERS
    from app.export import PROJECT_EXTRA, build_xlsx
    from app.store import dataset

    dataset.load_snapshot(finished_run.board.to_snapshot())
    ws = openpyxl.load_workbook(io.BytesIO(build_xlsx(finished_run.board.overlaps)))["projects"]
    head = [c.value for c in ws[1]]
    assert head[: len(PROJECT_HEADERS)] == PROJECT_HEADERS and head[len(PROJECT_HEADERS):] == PROJECT_EXTRA
    rows = [dict(zip(head, (c.value for c in r))) for r in ws.iter_rows(min_row=2)]
    sav = next(r for r in rows if r["filing_sponsor"] == "SAV")
    assert (sav["utility"], sav["sponsor"]) == ("Georgia Power", "GPC")
    assert {r["sponsor"] for r in rows if r["project_id"].startswith("DESC-")} == {"DESC"}


def test_api_lists_sources_for_the_ui(db):
    from app.main import app

    body = TestClient(app).get("/api/sources").json()
    desc, gpc = body["sources"]
    assert desc["display"]["short_name"] == "Dominion" and gpc["display"]["shape"] == "diamond"
    assert [s["code"] for s in gpc["sponsors"] if not s["default"]] == ["GTC", "MEAG", "DU"]
    assert body["bbox"] == [30.3, -85.7, 35.3, -78.4] and desc["builtin"]
