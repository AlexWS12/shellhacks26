# The Sources menu: plans people submit, read by their own Reader, compared as full peers.
import asyncio
import base64
import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import config
from app.clients.fetch import check_url
from app.core.htmltext import html_to_text, pages
from app.core.sheets import read_table, suggest, validate_mapping
from app.runtime.executor import execute
from app.runtime.run import new_run
from app.store import submissions

# Row 2 sits ~4 mi from Dominion's Jasper substation, so it must pair with Dominion and Georgia projects there.
CSV = (b"Project Name,Project ID,In-Service Date,From,To,Latitude,Longitude,State,Estimated Cost\n"
       b"Test Jasper Tap,TC-1,12/31/2027,,,32.40,-81.10,SC,\"$4,500,000\"\n"
       b"Test Line With No Date,TC-2,sometime,Alpha,Beta,,,SC,\n"
       b"Test Year Only,TC-3,2029,,,32.39,-81.12,SC,\n"
       b",TC-4,2028,,,,,,\n")
MAPPING = {"name": "Project Name", "id": "Project ID", "in_service": "In-Service Date", "endpoint_a": "From",
           "endpoint_b": "To", "lat": "Latitude", "lon": "Longitude", "state": "State", "cost": "Estimated Cost"}


@pytest.fixture
def plans_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "SUBMISSIONS_DIR", tmp_path / "subs")
    return tmp_path / "subs"


def test_store_saves_until_removed(plans_dir):
    s = submissions.create("Test Co", None, "SC", "spreadsheet", "../../evil name.csv", CSV, ".csv", columns=["a", "b"])
    assert s.status == "needs_mapping" and s.stored == f"{s.id}.csv" and s.filename == "evil name.csv"
    assert (plans_dir / s.stored).read_bytes() == CSV
    assert [x.id for x in submissions.list_all()] == [s.id]
    assert submissions.set_mapping(s.id, {"name": "a"}).status == "ready"
    assert submissions.delete(s.id) and submissions.list_all() == [] and not (plans_dir / s.stored).exists()
    assert submissions.get("../etc") is None


@pytest.mark.parametrize("owner", ["Georgia Power", "Dominion Energy South Carolina", "dominion energy sc", "SCE&G"])
def test_built_in_owners_are_refused(plans_dir, owner):
    with pytest.raises(ValueError, match="already built in"):
        submissions.create(owner, None, "SC", "spreadsheet", "x.csv", CSV, ".csv")


def test_columns_are_suggested_and_checked(tmp_path):
    p = tmp_path / "plan.csv"
    p.write_bytes(CSV)
    header, rows = read_table(p)
    assert header[:3] == ["Project Name", "Project ID", "In-Service Date"] and len(rows) == 4
    got = suggest(header)
    assert got["name"] == "Project Name" and got["in_service"] == "In-Service Date" and got["lat"] == "Latitude"
    with pytest.raises(ValueError, match="In-service"):
        validate_mapping({"name": "Project Name"}, header)
    with pytest.raises(ValueError, match="Not a column"):
        validate_mapping({"name": "Nope", "in_service": "In-Service Date"}, header)


def test_xlsx_with_a_title_row(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Some utility, 10-year plan"])
    ws.append(["Name", "Completion"])
    ws.append(["Test Sub", 2030])
    p = tmp_path / "plan.xlsx"
    wb.save(p)
    header, rows = read_table(p)
    assert header == ["Name", "Completion"] and rows == [(3, ["Test Sub", 2030])]  # real file row number


def _run():
    from app.pipeline import build_pipeline

    config.GEMINI_API_KEY, config.JEV_PROVIDER, config.OSM_LIVE = "", "", False
    run = new_run("live")
    run.pace = 0
    agents, sources = build_pipeline()
    asyncio.run(execute(run, agents, sources))
    return run


@pytest.fixture
def peer_run(plans_dir):
    s = submissions.create("Test Co", "Test Co", "SC", "spreadsheet", "plan.csv", CSV, ".csv",
                           columns=["Project Name", "Project ID", "In-Service Date", "From", "To", "Latitude", "Longitude",
                                    "State", "Estimated Cost"])
    submissions.set_mapping(s.id, MAPPING)
    return s, _run()


def test_submitted_plan_gets_its_own_reader_and_peers(peer_run):
    s, run = peer_run
    last = run.events[-1]
    assert last["type"] == "run.done" and last["ok"], last
    started = next(e for e in run.events if e["type"] == "run.started")
    assert f"extract_{s.id}" in [a["id"] for a in started["agents"]]
    assert s.id in [src["id"] for src in started["sources"]]
    geo = next(a for a in started["agents"] if a["id"] == "geocoder")
    assert f"extract_{s.id}" in geo["depends_on"]
    mine = [p for p in run.board.projects.values() if p.utility == s.owner_key]
    assert sorted(p.name for p in mine) == ["Test Jasper Tap", "Test Year Only"]
    tap = next(p for p in mine if p.name == "Test Jasper Tap")
    assert (tap.lat, tap.lon, tap.location_confidence, tap.cost_total) == (32.40, -81.10, "verified", 4500000)
    assert tap.endpoints[0].method == "submitted" and tap.source_ref == "row 2, ID TC-1"
    year = next(p for p in mine if p.name == "Test Year Only")
    assert (year.in_service_date, year.date_precision) == ("2029-12-31", "year")


def test_peer_pairs_join_the_ranking_but_not_the_benchmark(peer_run):
    s, run = peer_run
    b = run.board
    peer = [o for o in b.overlaps if b.projects[o.project_b].utility == s.owner_key]
    assert peer, "expected pairs with the submitted plan near Jasper"
    assert {b.projects[o.project_a].utility for o in peer} <= {"DESC", "GA"}
    assert all(o.distance_mi < 25 for o in b.overlaps)
    assert [o.rank for o in b.overlaps] == list(range(1, len(b.overlaps) + 1))
    # active pairs first, each group closest first
    assert [(o.finished, o.distance_mi) for o in b.overlaps] == sorted((o.finished, o.distance_mi) for o in b.overlaps)
    assert len(b.reference) == 6 and all(r.passed for r in b.reference)
    assert all(k.startswith("DESC-") and "|GA-" in k for k in b.analyses)  # written sides stay Dominion-Georgia
    assert all(o.id in b.costs for o in peer)
    assert "Test Co" in b.report["owners"]


def test_skipped_rows_are_reported(peer_run):
    s, run = peer_run
    c = next(c for c in run.board.checks if c.rule == "submitted_rows_skipped")
    assert "row 3: in-service date 'sometime' is not a date" in c.detail and "row 5: no project name" in c.detail
    assert any(c.rule == "submitted_dates_approx" for c in run.board.checks)


def test_export_uses_owner_names(peer_run):
    from app.export import build_xlsx
    from app.store import dataset

    s, run = peer_run
    dataset.load_snapshot(run.board.to_snapshot())
    wb = openpyxl.load_workbook(io.BytesIO(build_xlsx(run.board.overlaps)))
    owners = {r[6] for r in wb["overlaps"].iter_rows(min_row=2, values_only=True)}
    assert "Test Co" in owners
    assert "Test Co" in {r[1] for r in wb["projects"].iter_rows(min_row=2, values_only=True)}


def test_pdf_without_gemini_is_explained_not_fatal(plans_dir):
    s = submissions.create("Pdf Co", None, "GA", "pdf", "plan.pdf", b"%PDF-1.4 not really a pdf", ".pdf")
    run = _run()
    assert run.events[-1]["ok"], run.events[-1]
    done = next(e for e in run.events if e["type"] == "agent.done" and e["agent_id"] == f"extract_{s.id}")
    assert done["summary"].startswith(("not read", "stopped"))
    assert any(c.rule in ("submission_needs_gemini", "submission_unreadable") for c in run.board.checks)


def test_api_menu_flow(plans_dir):
    from app.main import app

    c = TestClient(app)
    r = c.post("/api/submissions", json={"owner": "Test Co", "state": "SC", "kind": "spreadsheet", "filename": "plan.csv",
                                         "content_b64": base64.b64encode(CSV).decode()})
    assert r.status_code == 200, r.text
    body = r.json()
    sid = body["submission"]["id"]
    assert body["submission"]["status"] == "needs_mapping" and body["preview"]["total_rows"] == 4
    assert body["suggested"]["in_service"] == "In-Service Date"
    assert c.put(f"/api/submissions/{sid}/mapping", json={"mapping": {"name": "Project Name"}}).status_code == 400
    ok = c.put(f"/api/submissions/{sid}/mapping", json={"mapping": MAPPING}).json()
    assert ok["usable"] == 2 and ok["submission"]["status"] == "ready" and len(ok["skipped"]) == 2
    listing = c.get("/api/submissions").json()
    assert [x["id"] for x in listing["submissions"]] == [sid] and listing["fields"]["name"]["required"]
    assert any(a["id"] == f"extract_{sid}" for a in c.get("/api/agents").json()["agents"])
    bad = c.post("/api/submissions", json={"owner": "X", "kind": "spreadsheet", "filename": "a.csv", "content_b64": "%%%"})
    assert bad.status_code == 400
    assert c.post("/api/submissions", json={"owner": "Georgia Power", "kind": "spreadsheet", "filename": "a.csv",
                                            "content_b64": base64.b64encode(CSV).decode()}).status_code == 400
    assert c.delete(f"/api/submissions/{sid}").json()["ok"] and c.get("/api/submissions").json()["submissions"] == []


@pytest.mark.parametrize("url", ["http://127.0.0.1/x", "http://localhost:8000/api", "http://10.1.2.3/plan.pdf",
                                 "http://169.254.169.254/latest/meta-data", "file:///etc/passwd", "ftp://example.org/x"])
def test_links_must_be_public(url):
    with pytest.raises(ValueError):
        check_url(url)


def test_web_page_text():
    html = "<html><head><style>x{}</style><script>bad()</script></head><body><h1>Plan</h1><table><tr><td>Line A</td><td>2028</td></tr></table></body></html>"
    text = html_to_text(html)
    assert "bad()" not in text and "Line A | 2028" in text  # a table row stays on one line
    assert pages("a\n\n" * 5) and len(pages(("word " * 1500 + "\n\n") * 3)) == 3
    rows = "\n".join(f"Line {i} | 2028" for i in range(900))  # one long table: split between rows, never inside one
    assert all(chunk.strip().split("\n\n")[-1].endswith("2028") for chunk in pages(rows))


def test_exports_store_formula_like_text_as_text():
    from app.export import safe_cell

    assert safe_cell("=HYPERLINK(\"http://x\")") == "'=HYPERLINK(\"http://x\")"
    assert safe_cell("+1") == "'+1" and safe_cell("@SUM(A1)") == "'@SUM(A1)" and safe_cell("-2") == "'-2"
    assert safe_cell("Jasper Tap") == "Jasper Tap" and safe_cell(-81.1) == -81.1 and safe_cell(None) is None


def test_results_mode_uses_the_graph_of_the_shown_run(peer_run):
    from app.main import app
    from app.pipeline import publish  # noqa: F401  (publish writes these two keys)
    from app.store import dataset

    s, run = peer_run
    started = next(e for e in run.events if e["type"] == "run.started")
    dataset.load_snapshot({**run.board.to_snapshot(), "agents": started["agents"], "sources": started["sources"]})
    submissions.delete(s.id)  # removed after the run: the shown run still had its Reader
    c = TestClient(app)
    latest = [a["id"] for a in c.get("/api/agents?of=latest").json()["agents"]]
    nxt = [a["id"] for a in c.get("/api/agents").json()["agents"]]
    assert f"extract_{s.id}" in latest and f"extract_{s.id}" not in nxt
    assert any(src.get("owner_key") == s.owner_key for src in started["sources"])


def test_row_numbers_match_the_file_with_blank_rows(tmp_path):
    p = tmp_path / "plan.csv"
    p.write_bytes(b"My plan\n\nName,In service\nA,2027\n\nB,2028\n")
    header, rows = read_table(p)
    assert header == ["Name", "In service"] and [n for n, _ in rows] == [4, 6]


@pytest.mark.parametrize("owner", ["Georgia Power Co.", "Dominion Energy, Inc.", "Dominion Energy South Carolina, Inc."])
def test_more_built_in_spellings_are_refused(plans_dir, owner):
    with pytest.raises(ValueError, match="already built in"):
        submissions.create(owner, None, "SC", "spreadsheet", "x.csv", CSV, ".csv")


def test_coordinates_outside_the_region_are_ignored(plans_dir):
    from app.core.sheets import to_project

    s = submissions.create("Test Co", None, "SC", "spreadsheet", "plan.csv", CSV, ".csv", columns=["a"])
    header = ["Name", "Date", "Lat", "Lon", "From", "To"]
    p, _ = to_project(s, {"name": "Name", "in_service": "Date", "lat": "Lat", "lon": "Lon", "endpoint_a": "From",
                          "endpoint_b": "To"}, header, ["X", "2027", "32.4", "81.1", "Alpha", "Beta"], 2)  # missing minus
    assert p and [e.name for e in p.endpoints] == ["Alpha", "Beta"] and all(e.lat is None for e in p.endpoints)


def test_spatial_buckets_find_exactly_the_same_pairs(finished_run):
    from app.core.models import Project
    from app.core.overlap import Filters, closest_mi, find_overlaps, geometry, span_ok, visible

    projects = list(finished_run.board.projects.values())
    # a synthetic submitted owner scattered over both overlap areas
    for k in range(40):
        projects.append(Project(id=f"peer-{k}", utility="own-peer", sponsor="Peer", name=f"P{k}", in_service_date="2029-12-31",
                                source_file="x", source_page=1, source_ref="r", lat=32.0 + (k % 8) * 0.22, lon=-82.4 + (k // 8) * 0.35,
                                location_confidence="verified"))
    f = Filters()
    got = {(o.project_a, o.project_b, o.distance_mi) for o in find_overlaps(projects, f)}
    shown = [p for p in projects if visible(p, f) and span_ok(p)]
    brute = set()
    for i, a in enumerate(shown):
        for b in shown[i + 1:]:
            d = closest_mi(geometry(a), geometry(b))
            if a.utility != b.utility and d < 25:
                x, y = sorted((a, b), key=lambda p: ({"DESC": 0, "GA": 1}.get(p.utility, 2), p.utility))
                brute.add((x.id, y.id, round(d, 2)))
    assert got == brute and len(got) > len(finished_run.board.overlaps)


def test_web_page_reader_stops_at_its_budget_and_the_run_still_publishes(plans_dir, monkeypatch):
    import types

    import app.agents.submissions as reader
    from app.clients import models

    calls = []

    async def fake_call(role, prompt, schema=None, *, system="", cache=True):
        calls.append(prompt)
        return models.Result({"projects": [{"name": f"Line {len(calls)}", "in_service": "2029"}]}, "gemini", "m1")

    # only the Reader sees this model; every other agent keeps Gemini off
    monkeypatch.setattr(reader, "models", types.SimpleNamespace(
        available=lambda role: True, call=fake_call, RoleExhausted=models.RoleExhausted))
    html = "".join(f"<p>{'filler ' * 900} Line {i} in service 2029</p>" for i in range(6)).encode()
    s = submissions.create("Web Co", None, "GA", "url", "plan.html", html, ".html", url="https://example.org/plan")
    monkeypatch.setattr(reader, "READER_BUDGET_S", -1.0)  # budget already used up: read nothing more
    run = _run()
    assert run.events[-1]["ok"], run.events[-1]
    assert calls == [] and any(c.rule == "submission_partial" for c in run.board.checks)
    monkeypatch.setattr(reader, "READER_BUDGET_S", 150.0)
    run = _run()
    mine = [p for p in run.board.projects.values() if p.utility == s.owner_key]
    assert len(calls) >= 2 and len(mine) == len(calls)
    assert all(p.source_file == "https://example.org/plan" and p.extracted_by == "gemini" for p in mine)
    assert len({p.id for p in mine}) == len(mine)
